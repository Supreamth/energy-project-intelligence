from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

wb = load_workbook("/tmp/justpow-plants.xlsx", data_only=True)
print("sheets", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("\nSHEET", name, "dims", ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column)
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    print("headers", headers)
    if ws.max_row < 2:
        continue
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    print("nrows", len(rows))
    print("sample", rows[0][:12] if rows else None)
    # fuel-like columns
    for i, h in enumerate(headers):
        if h is None:
            continue
        hs = str(h).lower()
        if any(k in hs for k in ["fuel", "เชื้อ", "type", "lat", "lon", "พิกัด", "ละติ", "ลองจิ", "solar", "แสง"]):
            vals = Counter(str(r[i])[:40] for r in rows if r[i] is not None)
            print(" col", h, "top", vals.most_common(8))
