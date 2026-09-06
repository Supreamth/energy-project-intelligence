from pathlib import Path

from openpyxl import load_workbook
import psycopg

wb = load_workbook("/tmp/justpow-plants.xlsx", data_only=True)
ws = wb["โรงไฟฟ้า"]
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
idx = {h: i for i, h in enumerate(headers)}
solar = []
for row in ws.iter_rows(min_row=2, values_only=True):
    fuel = str(row[idx["fuel type"]] or "")
    if fuel.lower() != "solar":
        continue
    solar.append(
        {
            "company": (row[idx["บริษัท"]] or "").strip() if row[idx["บริษัท"]] else "",
            "district": row[idx["อำเภอ"]],
            "province": row[idx["จังหวัด/ประเทศ"]],
            "mw": row[idx["กำลังการผลิตติดตั้ง(MW)"]],
        }
    )
print("solar_rows", len(solar), "unique_company", len({r["company"] for r in solar}))
with psycopg.connect("postgresql://hermes@127.0.0.1:55432/intelligence") as conn:
    names = {
        r[0]
        for r in conn.execute(
            "SELECT legal_name FROM intelligence.organizations"
        )
    }
hit = [r["company"] for r in solar if r["company"] in names]
print("exact_org_match_rows", len(hit), "unique", len(set(hit)))
print("sample solar", solar[:3])
print("has_lat_cols", False)
